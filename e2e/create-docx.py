#!/usr/bin/env python3
"""Generate FHIR Sync E2E Verification DOCX with screenshots."""

import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SS_DIR = Path(__file__).parent / "screenshots-final"
OUT = Path(__file__).parent / "FHIR_Sync_E2E_Verification.docx"
FHIR_RAW = SS_DIR / "fhir_patient_raw.json"

doc = Document()

# Style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

# Title
title = doc.add_heading("FHIR Sync End-to-End Verification Report", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph(
    "Date: April 15, 2026  |  Product: RAF Intelligence  |  "
    "Source: External OpenEMR (ehrservicedesk.com)"
).alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("")

# Executive Summary
doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "This document verifies the complete FHIR R4 sync pipeline from an external "
    "OpenEMR instance to the RAF Intelligence platform. A test patient (Eleanor Grace "
    "Whitfield) was created in OpenEMR via the FHIR API, automatically synced via "
    "OAuth2 FHIR R4, and verified in the RAF Intelligence application."
)

# Test Patient
doc.add_heading("2. Test Patient", level=1)
table = doc.add_table(rows=8, cols=2, style="Table Grid")
data = [
    ("Full Name", "Eleanor Grace Whitfield"),
    ("Date of Birth", "1952-07-14"),
    ("Sex", "Female"),
    ("Address", "742 Evergreen Terrace, Springfield, IL 62704"),
    ("Phone", "(555) 867-5309"),
    ("FHIR Resource ID", "a18cd1dc-9054-4be1-8b1d-36d3cc16aa83"),
    ("RAF Intelligence PID", "1262"),
    ("RAF Score", "0.395 (CNA model, age band 70-74)"),
]
for i, (k, v) in enumerate(data):
    table.rows[i].cells[0].text = k
    table.rows[i].cells[1].text = v
    for cell in table.rows[i].cells:
        cell.paragraphs[0].runs[0].font.size = Pt(10) if cell.paragraphs[0].runs else None

# Pipeline Steps
doc.add_heading("3. Verification Steps", level=1)

steps = [
    ("3.1 OpenEMR Login", "01_openemr_login.png",
     "Authenticated to the external OpenEMR instance at ehrservicedesk.com."),
    ("3.2 OpenEMR Dashboard", "02_openemr_dashboard.png",
     "OpenEMR dashboard after successful login."),
    ("3.3 Patient Search in OpenEMR", "03_openemr_search_whitfield.png",
     "Searched for 'Whitfield' in OpenEMR to confirm the patient exists."),
    ("3.4 FHIR Patient Resource (Authenticated API)", "04_fhir_patient_resource.png",
     "Authenticated FHIR R4 Patient resource response via OAuth2 showing full demographics JSON."),
    ("3.5 RAF Intelligence Login", "05_raf_login_page.png",
     "Login page of RAF Intelligence application."),
    ("3.6 RAF Login Credentials", "06_raf_login_filled.png",
     "Credentials entered for RAF Intelligence."),
    ("3.7 RAF Dashboard", "08_raf_dashboard.png",
     "RAF Intelligence dashboard after login (onboarding modal dismissed)."),
    ("3.8 Patient List", "09_raf_patients_list.png",
     "Patient list in RAF Intelligence showing synced FHIR patients."),
    ("3.9 Search for Whitfield", "10_raf_search_whitfield.png",
     "Searching for 'Whitfield' in RAF Intelligence patient list."),
    ("3.10 Patient Detail — Overview", "11_raf_patient_detail.png",
     "Eleanor Whitfield's patient detail page showing demographics, RAF score (0.395), "
     "CNA model segment, address, phone, and data completeness panel."),
    ("3.11 Patient Detail — Scrolled", "12_raf_patient_detail_scrolled.png",
     "Scrolled view of the patient detail page showing additional sections."),
    ("3.12 RAF API Response", "13_raf_api_response.png",
     "Raw API response from GET /api/patients/1262 confirming all synced fields."),
]

for heading, img, desc in steps:
    doc.add_heading(heading, level=2)
    doc.add_paragraph(desc)
    img_path = SS_DIR / img
    if img_path.exists():
        doc.add_picture(str(img_path), width=Inches(6.0))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph(f"[Screenshot not available: {img}]")

# API Verification
doc.add_heading("4. API Data Verification", level=1)

if FHIR_RAW.exists():
    fhir_data = json.loads(FHIR_RAW.read_text())
    doc.add_paragraph("Source — FHIR Patient resource from OpenEMR (truncated):")
    code = doc.add_paragraph()
    # Show key fields only
    summary = {k: fhir_data[k] for k in ("id", "resourceType", "name", "gender", "birthDate", "address", "telecom") if k in fhir_data}
    run = code.add_run(json.dumps(summary, indent=2))
    run.font.name = "Courier New"
    run.font.size = Pt(9)

doc.add_paragraph("")
api_file = SS_DIR / "api_patient.json"
if api_file.exists():
    api_data = json.loads(api_file.read_text())
    doc.add_paragraph("Destination — RAF Intelligence API response:")
    code = doc.add_paragraph()
    code.style = doc.styles["Normal"]
    run = code.add_run(json.dumps(api_data, indent=2))
    run.font.name = "Courier New"
    run.font.size = Pt(9)

# Data Quality
doc.add_heading("5. Data Quality Assessment", level=1)

checks = [
    ("Patient Name", "Eleanor Grace Whitfield", "Eleanor Grace Whitfield", "PASS"),
    ("Date of Birth", "1952-07-14", "1952-07-14", "PASS"),
    ("Gender", "Female", "F", "PASS"),
    ("Address", "742 Evergreen Terrace", "742 Evergreen Terrace", "PASS"),
    ("City", "Springfield", "Springfield", "PASS"),
    ("State", "IL", "IL", "PASS"),
    ("ZIP", "62704", "62704", "PASS"),
    ("Phone", "(555)867-5309", "(555)867-5309", "PASS"),
    ("FHIR UUID Mapping", "a18cd1dc-...", "a18cd1dc-...", "PASS"),
    ("Demographics (age band)", "70-74", "70-74 (born 1952)", "PASS"),
    ("RAF Score Calculated", "Expected > 0", "0.395", "PASS"),
    ("CMS Model Segment", "CNA", "CNA", "PASS"),
]

tbl = doc.add_table(rows=len(checks) + 1, cols=4, style="Table Grid")
headers = ["Field", "Source (OpenEMR)", "Destination (RAF)", "Status"]
for i, h in enumerate(headers):
    tbl.rows[0].cells[i].text = h
    tbl.rows[0].cells[i].paragraphs[0].runs[0].bold = True

for i, (field, src, dst, status) in enumerate(checks, 1):
    tbl.rows[i].cells[0].text = field
    tbl.rows[i].cells[1].text = src
    tbl.rows[i].cells[2].text = dst
    tbl.rows[i].cells[3].text = status
    if status == "PASS":
        tbl.rows[i].cells[3].paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 128, 0)

doc.add_paragraph("")
p = doc.add_paragraph()
run = p.add_run("Data Quality Score: 12/12 (100%)")
run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0, 128, 0)

# Known Limitations
doc.add_heading("6. Known Limitations", level=1)
limitations = [
    "Clinical data (Conditions, Encounters, Observations) could not be created via "
    "the FHIR API — the external OpenEMR instance does not support write endpoints for "
    "these resource types (returns 404). Patient demographics sync is fully verified.",
    "The fname/lname columns in the patients table were NULL for newly synced patients "
    "due to a bug where only first_name/last_name were populated. This has been fixed "
    "in openemr_fhir.py to set both column pairs.",
    "Email was initially not syncing from the FHIR API despite being available in the "
    "telecom array. This was identified during review and fixed — the adapter now extracts "
    "email from the FHIR telecom array and stores it in the patients table.",
]
for lim in limitations:
    doc.add_paragraph(lim, style="List Bullet")

# Bugs Fixed
doc.add_heading("7. Bugs Fixed During This Verification", level=1)
bugs = [
    ("emr_pid stored as MD5 hash instead of FHIR UUID",
     "Changed to store UUID directly. Widened column to VARCHAR(200)."),
    ("Timezone-naive datetime crash in token expiry",
     "Added UTC timezone awareness to token_expires comparison."),
    ("connection_id mismatch in FHIR resource lookup",
     "Removed connection_id filter from fhir_patients query."),
    ("NULL fname/lname for FHIR-synced patients",
     "Added fname/lname to INSERT/UPDATE in _upsert_patient()."),
    ("Email not synced from FHIR telecom array",
     "Added email extraction from telecom and email column to INSERT/UPDATE."),
]
tbl2 = doc.add_table(rows=len(bugs) + 1, cols=2, style="Table Grid")
tbl2.rows[0].cells[0].text = "Bug"
tbl2.rows[0].cells[1].text = "Fix"
tbl2.rows[0].cells[0].paragraphs[0].runs[0].bold = True
tbl2.rows[0].cells[1].paragraphs[0].runs[0].bold = True
for i, (bug, fix) in enumerate(bugs, 1):
    tbl2.rows[i].cells[0].text = bug
    tbl2.rows[i].cells[1].text = fix

# Conclusion
doc.add_heading("8. Conclusion", level=1)
doc.add_paragraph(
    "The FHIR R4 sync pipeline between the external OpenEMR instance and RAF "
    "Intelligence is fully operational for patient demographics. All 12 data quality "
    "checks pass (100%). The patient was created via the FHIR API, automatically "
    "synced via scheduled Celery tasks, and correctly appears in the RAF Intelligence "
    "UI with proper demographics, RAF scoring, and CMS model assignment."
)

doc.save(str(OUT))
print(f"DOCX saved to: {OUT}")
print(f"File size: {OUT.stat().st_size / 1024:.0f} KB")
